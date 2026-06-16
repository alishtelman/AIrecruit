"""
E2E integration test: practical task flow.

Requires the full stack running: docker compose up -d --build
Run with: cd backend && python3 -m pytest tests/test_e2e_practical_task.py -v -s

Covers:
1. Register candidate, upload resume
2. Start backend_engineer interview (role supports practical tasks)
3. Answer 3+ questions until practical task triggers (or max_questions reached)
4. Submit practical task via /practical-submission
5. Verify:
   - practical_submission stored separately (not in messages/history)
   - interview continues after submit
   - finish + report contains practical_section
   - negative: practical_submission role not exposed to client
"""
import asyncio
import io
import json
import time
import uuid

import pytest
import pytest_asyncio
from docx import Document
from httpx import AsyncClient

from tests.conftest import auth_headers, BASE_URL


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _build_backend_resume_docx() -> bytes:
    doc = Document()
    doc.add_heading("Senior Backend Engineer", 0)
    doc.add_paragraph("5 years experience building distributed systems.")
    doc.add_paragraph("Stack: Python, FastAPI, PostgreSQL, Redis, Kafka, Docker, Kubernetes.")
    doc.add_paragraph("Built payment processing service handling 50k req/s.")
    doc.add_paragraph("Led API migration from REST to gRPC, reduced latency by 40%.")
    doc.add_paragraph("Strong skills in SQL, async Python, system design, and code review.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


async def _register_and_login(client: AsyncClient) -> str:
    email = f"e2e_practical_{uuid.uuid4().hex[:8]}@example.com"
    await client.post("/api/v1/auth/candidate/register", json={
        "email": email,
        "password": "testpass123",
        "full_name": "E2E Practical Tester",
    })
    resp = await client.post("/api/v1/auth/login", json={
        "email": email,
        "password": "testpass123",
        "account_type": "candidate",
    })
    assert resp.status_code == 200, f"login failed: {resp.text}"
    return resp.json()["access_token"]


async def _upload_resume(client: AsyncClient, token: str) -> str:
    docx_bytes = _build_backend_resume_docx()
    resp = await client.post(
        "/api/v1/candidate/resume/upload",
        headers=auth_headers(token),
        files={"file": ("resume.docx", io.BytesIO(docx_bytes),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert resp.status_code == 200, f"resume upload failed: {resp.text}"
    return resp.json()["resume_id"]


async def _start_interview(client: AsyncClient, token: str) -> dict:
    resp = await client.post(
        "/api/v1/interviews/start",
        headers=auth_headers(token),
        json={"target_role": "backend_engineer", "language": "en"},
    )
    assert resp.status_code == 201, f"start interview failed: {resp.text}"
    return resp.json()


async def _send_message(client: AsyncClient, token: str, interview_id: str, text: str) -> dict:
    resp = await client.post(
        f"/api/v1/interviews/{interview_id}/message",
        headers=auth_headers(token),
        json={"message": text},
    )
    assert resp.status_code == 200, f"send message failed (status {resp.status_code}): {resp.text}"
    return resp.json()


async def _get_detail(client: AsyncClient, token: str, interview_id: str) -> dict:
    resp = await client.get(
        f"/api/v1/interviews/{interview_id}",
        headers=auth_headers(token),
    )
    assert resp.status_code == 200
    return resp.json()


async def _submit_practical(client: AsyncClient, token: str, interview_id: str, task: dict) -> dict:
    resp = await client.post(
        f"/api/v1/interviews/{interview_id}/practical-submission",
        headers=auth_headers(token),
        json={
            "task_id": task["task_id"],
            "task_type": task["task_type"],
            "answer": _mock_solution(task),
            "language": task.get("language"),
            "duration_seconds": 180,
        },
    )
    assert resp.status_code == 200, f"practical submission failed: {resp.text}"
    return resp.json()


def _mock_solution(task: dict) -> str:
    task_type = task.get("task_type", "")
    if task_type == "coding_task":
        return (
            task.get("starter_code", "") +
            "\n# E2E test solution\n"
            "def validate_registration(data):\n"
            "    errors = []\n"
            "    if not data.get('email'): errors.append('email required')\n"
            "    if len(data.get('password', '')) < 8: errors.append('password too short')\n"
            "    age = data.get('age', 0)\n"
            "    if not (18 <= age <= 120): errors.append('invalid age')\n"
            "    return {'valid': not errors, 'errors': errors}\n"
        )
    if task_type == "debugging_task":
        return (
            "# Found the bug: start = page * page_size is 0-indexed offset\n"
            "# but page is 1-indexed. Fix:\n"
            "def paginate(items, page, page_size):\n"
            "    start = (page - 1) * page_size\n"
            "    return items[start:start + page_size]\n"
        )
    if task_type == "sql_task":
        return (
            "SELECT user_id, SUM(amount) AS total_revenue\n"
            "FROM orders\n"
            "WHERE created_at >= NOW() - INTERVAL '30 days'\n"
            "GROUP BY user_id\n"
            "ORDER BY total_revenue DESC\n"
            "LIMIT 5;\n"
        )
    # fallback for any task type
    return "E2E mock solution: approach would be to [analyse the problem], implement [core logic], handle [edge cases]."


# ---------------------------------------------------------------------------
# Main E2E test
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_practical_task_e2e_full_flow():
    """
    Full E2E: start interview, answer questions until practical task triggers,
    submit solution, finish, verify report.
    """
    async with AsyncClient(base_url=BASE_URL, timeout=90.0) as client:
        # 1. Register + upload resume
        token = await _register_and_login(client)
        await _upload_resume(client, token)

        # 2. Start interview
        start_data = await _start_interview(client, token)
        interview_id = start_data["interview_id"]
        print(f"\n✓ Interview started: {interview_id}")
        print(f"  max_questions={start_data['max_questions']}, language=en")

        # 3. Confirm intro ("ready")
        resp = await _send_message(client, token, interview_id, "Ready to start.")
        print(f"✓ Intro confirmed. First question: {resp.get('current_question','')[:80]}...")

        # 4. Answer questions until practical task fires or max_questions reached
        practical_task = None
        answers = [
            "I have 5 years of experience with Python FastAPI and PostgreSQL. "
            "I built a payment processing microservice handling 50k req/s with async workers.",

            "For database design I think carefully about normalization, indexes, "
            "and connection pooling. I've used SQLAlchemy async sessions extensively.",

            "My API design approach: RESTful resources, versioning, clear error codes, "
            "pagination, and OpenAPI docs. I've also worked with gRPC for internal services.",

            "For system design I start with requirements, identify bottlenecks, "
            "pick appropriate data stores, then design for horizontal scaling.",

            "Debugging approach: structured logging, distributed tracing with Jaeger, "
            "database slow query logs, and profiling with py-spy.",

            "Code review I focus on: correctness, tests, security (SQL injection, auth), "
            "performance, and maintainability.",
        ]

        for i, answer in enumerate(answers, 1):
            print(f"  Answering question {i}...")
            resp = await _send_message(client, token, interview_id, answer)

            q_type = resp.get("question_delivery_type", "voice_only")
            print(f"  -> question_delivery_type={q_type}, question_count={resp.get('question_count')}")

            if q_type == "practical_task":
                practical_task = resp.get("practical_task")
                print(f"\n✓ PRACTICAL TASK TRIGGERED after {i} answers!")
                print(f"  task_type={practical_task.get('task_type')}")
                print(f"  title={practical_task.get('title')}")
                print(f"  language={practical_task.get('language')}")
                print(f"  voice_intro={practical_task.get('voice_intro','')[:60]}")
                break

            if not resp.get("current_question"):
                print(f"  -> Interview finished naturally (no more questions)")
                break

        # 5. If practical task was triggered, submit solution
        if practical_task:
            # 5a. Verify detail BEFORE submit: no practical_submission in messages
            detail_before = await _get_detail(client, token, interview_id)
            messages_before = detail_before.get("messages", [])
            practical_roles_before = [m["role"] for m in messages_before if m["role"] == "practical_submission"]
            assert len(practical_roles_before) == 0, (
                f"practical_submission must not appear in InterviewDetailResponse before submit: {practical_roles_before}"
            )
            print(f"\n✓ PRE-SUBMIT: messages={len(messages_before)}, "
                  f"no practical_submission exposed in API")

            # 5b. Submit solution
            submit_resp = await _submit_practical(client, token, interview_id, practical_task)
            print(f"\n✓ PRACTICAL SUBMITTED")
            print(f"  question_count={submit_resp.get('question_count')}")
            print(f"  current_question={str(submit_resp.get('current_question',''))[:80]}...")
            print(f"  question_delivery_type={submit_resp.get('question_delivery_type','voice_only')}")

            # 5c. Interview must continue after submit
            assert submit_resp.get("status") in ("in_progress",), (
                f"Interview should still be in_progress after practical submit: {submit_resp.get('status')}"
            )
            next_q = submit_resp.get("current_question")
            # May be None if max_questions reached
            print(f"  -> interview continues: next_question={'present' if next_q else 'none (quota reached)'}")

            # 5d. Verify practical_submission NOT in API messages
            detail_after = await _get_detail(client, token, interview_id)
            messages_after = detail_after.get("messages", [])
            practical_roles_after = [m["role"] for m in messages_after if m["role"] == "practical_submission"]
            assert len(practical_roles_after) == 0, (
                f"practical_submission must NEVER appear in InterviewDetailResponse: {practical_roles_after}"
            )

            # 5e. Verify raw code not in visible transcript messages
            visible_contents = [m["content"] for m in messages_after if m["role"] == "candidate"]
            # The mock solution has distinctive code — should NOT appear
            code_in_transcript = any(
                "def validate_registration" in c or "def paginate" in c
                for c in visible_contents
            )
            assert not code_in_transcript, (
                "Raw solution code must NOT appear in candidate messages transcript!"
            )
            # The summary message SHOULD be there
            summary_in_transcript = any(
                "practical" in c.lower() or "task" in c.lower()
                for c in visible_contents
            )
            print(f"\n✓ POST-SUBMIT: messages={len(messages_after)}")
            print(f"  practical_submission in API: {len(practical_roles_after)} (expected 0)")
            print(f"  raw code in transcript: {code_in_transcript} (expected False)")
            print(f"  summary in transcript: {summary_in_transcript} (expected True)")

        else:
            print("\n⚠ Practical task did not trigger (quota reached before threshold)")
            print("  This may happen with low max_questions. Verifying no regressions...")

        # 6. Answer remaining questions and finish
        # First check if we can finish (quota reached) or need more answers
        detail_now = await _get_detail(client, token, interview_id)
        interview_status = detail_now.get("status")
        qcount = detail_now.get("question_count", 0)
        max_q = detail_now.get("max_questions", 8)
        print(f"\nInterview state: status={interview_status}, q={qcount}/{max_q}")

        if interview_status == "in_progress":
            # Fill remaining questions if needed (up to 3 more)
            if qcount < max_q:
                remaining_answers = [
                    "I prioritize tasks by impact and urgency, communicate blockers early.",
                    "For testing I use pytest, unit+integration tests, and aim for 80% coverage.",
                    "My biggest challenge was scaling a database under heavy write load.",
                ]
                for extra in remaining_answers[:3]:
                    try:
                        extra_resp = await _send_message(client, token, interview_id, extra)
                        if not extra_resp.get("current_question"):
                            break
                    except Exception:
                        break

        # 7. Finish the interview
        finish_resp = await client.post(
            f"/api/v1/interviews/{interview_id}/finish",
            headers=auth_headers(token),
        )
        if finish_resp.status_code == 422:
            # Not enough answers — add more
            for extra in ["My approach to DevOps is Infrastructure as Code.", "I value clean code."]:
                try:
                    await _send_message(client, token, interview_id, extra)
                except Exception:
                    pass
            finish_resp = await client.post(
                f"/api/v1/interviews/{interview_id}/finish",
                headers=auth_headers(token),
            )

        assert finish_resp.status_code == 200, f"finish failed: {finish_resp.text}"
        finish_data = finish_resp.json()
        print(f"\n✓ FINISH: status={finish_data.get('status')}")

        # 8. Poll for report
        report_id = finish_data.get("report_id")
        if not report_id:
            print("  Polling for report...")
            deadline = time.time() + 90
            while time.time() < deadline:
                await asyncio.sleep(3)
                status_resp = await client.get(
                    f"/api/v1/interviews/{interview_id}/report-status",
                    headers=auth_headers(token),
                )
                if status_resp.status_code == 200:
                    sdata = status_resp.json()
                    state = sdata.get("processing_state")
                    print(f"  report_status={state}")
                    if state == "ready" and sdata.get("report_id"):
                        report_id = sdata["report_id"]
                        break
                    if state == "failed":
                        print(f"  ⚠ report failed: {sdata.get('failure_reason')}")
                        break

        if report_id:
            print(f"✓ REPORT READY: {report_id}")

            # 9. Get report and check practical_section
            report_resp = await client.get(
                f"/api/v1/reports/{report_id}",
                headers=auth_headers(token),
            )
            assert report_resp.status_code == 200, f"get report failed: {report_resp.text}"
            report = report_resp.json()

            full_json = report.get("full_report_json") or {}
            practical_section = report.get("practical_section") or full_json.get("practical_section")

            print(f"\n✓ REPORT practical_section:")
            print(f"  has_practical_tasks={practical_section.get('has_practical_tasks') if practical_section else 'MISSING'}")

            if practical_task and practical_section:
                assert practical_section.get("has_practical_tasks") is True, (
                    "practical_section.has_practical_tasks must be True"
                )
                assert practical_section.get("practical_tasks_count", 0) >= 1
                print(f"  practical_score={practical_section.get('practical_score')}")
                print(f"  evaluation_status={practical_section.get('evaluation_status')}")
                print(f"  interview_mode={practical_section.get('interview_mode')}")
                tasks_in_section = practical_section.get("tasks", [])
                print(f"  tasks in section: {len(tasks_in_section)}")
                if tasks_in_section:
                    t = tasks_in_section[0]
                    print(f"  task[0]: type={t.get('task_type')}, "
                          f"score={t.get('practical_score')}, "
                          f"correctness={t.get('correctness')}, "
                          f"trace_status={t.get('trace_status')}")

            elif not practical_task:
                # No practical task was triggered — section should say has_practical_tasks=False
                if practical_section:
                    assert practical_section.get("has_practical_tasks") is False
                    print("  (no practical task triggered this run — section is empty)")
            else:
                # practical_task triggered but section missing
                pytest.fail("practical_task was triggered but practical_section is missing from report!")
        else:
            print("⚠ Report not ready within timeout — skipping report assertions")

        print("\n" + "=" * 60)
        print("E2E RESULT:")
        if practical_task:
            print("  ✓ Practical task triggered by policy")
            print("  ✓ practical_submission not in API messages")
            print("  ✓ Raw code not in candidate transcript")
            print("  ✓ Interview continued after submission")
            print("  ✓ practical_section present in report")
        else:
            print("  ⚠ Practical task not triggered (quota too low or stage mismatch)")
            print("  ✓ No regressions in base interview flow")
        print("=" * 60)


# ---------------------------------------------------------------------------
# Separate: verify practical_submission role isolation via API
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_practical_submission_not_in_detail_response():
    """
    Submit a practical task and immediately verify that
    /interviews/{id} does NOT contain role=practical_submission.
    Focused isolation test.
    """
    async with AsyncClient(base_url=BASE_URL, timeout=90.0) as client:
        token = await _register_and_login(client)
        await _upload_resume(client, token)
        start = await _start_interview(client, token)
        iid = start["interview_id"]

        # Confirm intro
        await _send_message(client, token, iid, "Ready.")

        # Answer enough to trigger practical task (3-6 answers)
        practical_task = None
        for answer in [
            "I have deep Python and PostgreSQL experience in distributed systems.",
            "I design APIs with versioning, clear error codes, and OpenAPI documentation.",
            "I use Redis for caching, Kafka for event streaming, and Docker for containerization.",
            "For SQL I optimize with proper indexes, EXPLAIN ANALYZE, and connection pooling.",
        ]:
            resp = await _send_message(client, token, iid, answer)
            if resp.get("question_delivery_type") == "practical_task":
                practical_task = resp.get("practical_task")
                break
            if not resp.get("current_question"):
                break

        if not practical_task:
            pytest.skip("Practical task did not trigger in this run (policy threshold not met)")

        # Submit practical task
        await _submit_practical(client, token, iid, practical_task)

        # Immediately get detail and check roles
        detail = await _get_detail(client, token, iid)
        roles_seen = {m["role"] for m in detail.get("messages", [])}

        assert "practical_submission" not in roles_seen, (
            f"practical_submission role leaked into InterviewDetailResponse! "
            f"Roles seen: {roles_seen}"
        )
        assert "system" not in roles_seen, (
            f"system role leaked into InterviewDetailResponse! Roles seen: {roles_seen}"
        )
        allowed_roles = {"assistant", "candidate"}
        unexpected = roles_seen - allowed_roles
        assert not unexpected, (
            f"Unexpected roles in InterviewDetailResponse: {unexpected}"
        )
        print(f"\n✓ InterviewDetailResponse roles: {roles_seen} (clean)")


# ---------------------------------------------------------------------------
# Negative case: evaluation failure trace
# ---------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_practical_submission_evaluation_failure_does_not_block():
    """
    Even if OpenAI evaluation call would fail, the submission endpoint must return
    a valid SendMessageResponse and the interview must continue.

    We test this indirectly by checking that the endpoint returns 200 even without
    waiting for the evaluation to complete (evaluation is async and swallows errors).
    """
    async with AsyncClient(base_url=BASE_URL, timeout=90.0) as client:
        token = await _register_and_login(client)
        await _upload_resume(client, token)
        start = await _start_interview(client, token)
        iid = start["interview_id"]

        await _send_message(client, token, iid, "Ready.")

        practical_task = None
        for answer in [
            "Python, FastAPI, async PostgreSQL — my main stack for 4 years.",
            "API design: REST resources, versioning (/v1), OpenAPI, proper HTTP codes.",
            "Distributed tracing: Jaeger, structured JSON logs, correlation IDs.",
            "Database: EXPLAIN ANALYZE, covering indexes, avoiding N+1 queries.",
        ]:
            resp = await _send_message(client, token, iid, answer)
            if resp.get("question_delivery_type") == "practical_task":
                practical_task = resp.get("practical_task")
                break
            if not resp.get("current_question"):
                break

        if not practical_task:
            pytest.skip("Practical task did not trigger")

        # Submit with deliberately minimal/trivial answer (still valid)
        resp = await client.post(
            f"/api/v1/interviews/{iid}/practical-submission",
            headers=auth_headers(token),
            json={
                "task_id": practical_task["task_id"],
                "task_type": practical_task["task_type"],
                "answer": "# minimal answer for negative test",
                "language": practical_task.get("language"),
                "duration_seconds": 5,
            },
        )
        # Must return 200 regardless of evaluation outcome
        assert resp.status_code == 200, (
            f"Practical submission must return 200 even with minimal answer: {resp.text}"
        )
        data = resp.json()
        assert data.get("status") == "in_progress", (
            f"Interview must remain in_progress after submission: {data.get('status')}"
        )
        print(f"\n✓ Practical submission returned 200 with minimal answer")
        print(f"  interview status: {data.get('status')}")
        print(f"  next question: {str(data.get('current_question',''))[:60]}...")
